from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path, PurePosixPath
from typing import Any

import numpy as np
from netCDF4 import Dataset, Group
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_SOURCE = ROOT / "external_cache" / "v0.3.0-field-18731994"
EXPECTED = {
    "LICENSE.txt": (1033, "4b5dac56b6af4b437f77f2cd1757c6d8"),
    "README.md": (4486, "78b1f3923739af43156b5eccc99cd391"),
    "TestMatrix.xlsx": (14113, "e5a4ab27620b89e01f4a5eb953cbe026"),
    "Multiple Wake.zip": (607488674, "fe1ccca62ad418f682b0dfae4f97a515"),
}


def digest(path: Path, algorithm: str) -> str:
    checksum = hashlib.new(algorithm)
    with path.open("rb") as handle:
        while chunk := handle.read(1024 * 1024):
            checksum.update(chunk)
    return checksum.hexdigest()


def safe_member(info: zipfile.ZipInfo) -> bool:
    path = PurePosixPath(info.filename.replace("\\", "/"))
    mode = info.external_attr >> 16
    return (
        bool(path.parts)
        and not path.is_absolute()
        and ".." not in path.parts
        and not (path.parts[0].endswith(":") if path.parts else False)
        and not ((mode & 0o170000) == 0o120000)
    )


def extract_safely(archive: Path, destination: Path) -> list[dict[str, Any]]:
    destination.mkdir(parents=True, exist_ok=True)
    destination_root = destination.resolve()
    inventory: list[dict[str, Any]] = []
    seen: set[str] = set()
    with zipfile.ZipFile(archive) as bundle:
        if bundle.testzip() is not None:
            raise SystemExit("Archive CRC verification failed")
        for info in bundle.infolist():
            normalized = str(PurePosixPath(info.filename.replace("\\", "/")))
            if not safe_member(info):
                raise SystemExit(f"Unsafe archive member: {info.filename}")
            if normalized.casefold() in seen:
                raise SystemExit(f"Duplicate archive member: {info.filename}")
            seen.add(normalized.casefold())
            target = (destination / Path(*PurePosixPath(normalized).parts)).resolve()
            if destination_root not in target.parents and target != destination_root:
                raise SystemExit(f"Archive member escapes destination: {info.filename}")
            row: dict[str, Any] = {
                "member": normalized,
                "compressed_size": info.compress_size,
                "uncompressed_size": info.file_size,
                "crc32": f"{info.CRC:08x}",
                "compression": info.compress_type,
                "is_directory": info.is_dir(),
            }
            if info.is_dir():
                target.mkdir(parents=True, exist_ok=True)
            else:
                target.parent.mkdir(parents=True, exist_ok=True)
                member_hash = hashlib.sha256()
                if target.exists() and target.stat().st_size == info.file_size:
                    row["extraction"] = "reused_existing_exact_size_then_rehashed"
                    row["sha256"] = digest(target, "sha256")
                else:
                    if target.exists():
                        raise SystemExit(f"Refusing to overwrite mismatched member: {target}")
                    with bundle.open(info) as source, target.open("xb") as sink:
                        while chunk := source.read(1024 * 1024):
                            member_hash.update(chunk)
                            sink.write(chunk)
                    row["extraction"] = "extracted"
                    row["sha256"] = member_hash.hexdigest()
            inventory.append(row)
    return inventory


def json_value(value: Any) -> Any:
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return value


def inspect_group(group: Group, prefix: str = "") -> dict[str, Any]:
    variables = []
    for name, variable in sorted(group.variables.items()):
        variables.append(
            {
                "path": f"{prefix}/{name}" or "/",
                "dtype": str(variable.dtype),
                "dimensions": list(variable.dimensions),
                "shape": list(variable.shape),
                "attributes": {
                    key: json_value(variable.getncattr(key)) for key in sorted(variable.ncattrs())
                },
            }
        )
    groups = [
        inspect_group(child, f"{prefix}/{name}")
        for name, child in sorted(group.groups.items())
    ]
    return {
        "path": prefix or "/",
        "dimensions": {
            name: {"size": len(dimension), "unlimited": dimension.isunlimited()}
            for name, dimension in sorted(group.dimensions.items())
        },
        "attributes": {key: json_value(group.getncattr(key)) for key in sorted(group.ncattrs())},
        "variables": variables,
        "groups": groups,
    }


def inspect_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        values = list(sheet.values)
        headers = [str(value) if value is not None else None for value in values[0]]
        rows = []
        for raw in values[1:]:
            if all(value is None for value in raw):
                continue
            rows.append(
                {header: json_value(value) for header, value in zip(headers, raw) if header}
            )
        sheets.append(
            {
                "title": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "headers": headers,
                "rows": rows,
            }
        )
    return {"sheets": sheets}


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Inspect frozen v0.3.0 field source without scoring"
    )
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    source = args.source_dir.resolve()
    custody = []
    for name, (expected_size, expected_md5) in EXPECTED.items():
        path = source / name
        if not path.is_file():
            raise SystemExit(f"Missing frozen source file: {path}")
        actual_size = path.stat().st_size
        actual_md5 = digest(path, "md5")
        if actual_size != expected_size or actual_md5 != expected_md5:
            raise SystemExit(f"Frozen source mismatch: {name}")
        custody.append(
            {
                "name": name,
                "size_bytes": actual_size,
                "md5": actual_md5,
                "sha256": digest(path, "sha256"),
            }
        )

    extraction = source / "extracted"
    inventory = extract_safely(source / "Multiple Wake.zip", extraction)
    netcdf_files = sorted(extraction.rglob("*.nc"))
    if not netcdf_files:
        raise SystemExit("No NetCDF members found")
    structures = []
    for path in netcdf_files:
        with Dataset(path, "r") as dataset:
            structures.append(
                {
                    "member": path.relative_to(extraction).as_posix(),
                    "sha256": digest(path, "sha256"),
                    "size_bytes": path.stat().st_size,
                    "format": dataset.file_format,
                    "structure": inspect_group(dataset),
                }
            )
    result = {
        "schema_version": "1.0.0",
        "inspection_mode": "STRUCTURE_AND_DOMAIN_METADATA_ONLY_NO_CLAIM_METRICS",
        "source_custody": custody,
        "archive_inventory": inventory,
        "test_matrix": inspect_workbook(source / "TestMatrix.xlsx"),
        "netcdf": structures,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    with args.output.open("w", encoding="utf-8", newline="\n") as handle:
        handle.write(json.dumps(result, indent=2, sort_keys=True) + "\n")


if __name__ == "__main__":
    main()
